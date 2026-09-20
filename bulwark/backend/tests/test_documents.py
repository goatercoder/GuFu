"""System security plan, POA&M export and the readiness report."""

from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient

from bulwark import catalog

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def populated(client: TestClient, system: dict) -> dict:
    """A system with a narrative, evidence, a POA&M item and an agent report."""
    sid = system["id"]
    client.put(
        f"/api/systems/{sid}",
        json={
            "boundary_description": "VLAN 20 behind the FortiGate.",
            "cui_description": "Drawings received under DFARS 252.204-7012.",
            "data_flow_description": "Portal -> SharePoint -> workstation -> CNC.",
        },
    )
    client.put(
        f"/api/systems/{sid}/controls/3.1.1",
        json={"status": "implemented", "implementation_narrative": "Unique accounts in Entra ID."},
    )
    client.put(f"/api/systems/{sid}/controls/3.13.14",
               json={"status": "not_applicable", "na_justification": "No VoIP is used."})
    client.post(f"/api/systems/{sid}/assets",
                json={"name": "SHOP-CNC-01", "asset_type": "iot_ot", "category": "specialized",
                      "category_rationale": "Machine controller on an isolated VLAN."})
    client.post(f"/api/systems/{sid}/assets",
                json={"name": "MKT-LAPTOP", "category": "out_of_scope",
                      "category_rationale": "Guest VLAN with no route to the enclave."})
    client.post(f"/api/systems/{sid}/evidence",
                data={"title": "Access Control Policy", "kind": "policy", "control_ids": ["3.1.1"]},
                files={"file": ("acp.md", io.BytesIO(b"# policy"), "text/markdown")})
    client.post(f"/api/systems/{sid}/poam",
                json={"control_id": "3.5.3", "title": "Roll out MFA",
                      "milestones": [{"description": "Buy keys"}]})
    key = client.get(f"/api/systems/{sid}/enrollment").json()["enrollment_key"]
    client.post(
        "/api/agents/report",
        headers={"Authorization": f"Bearer {key}"},
        json={
            "schema_version": 1,
            "agent": {"version": "0.1.0", "platform": "windows"},
            "asset": {"hostname": "SHOP-CNC-01", "os_family": "windows", "os_name": "Windows 10"},
            "checks": [{"check_id": "os.disk.encryption", "status": "fail", "observed": "not encrypted"}],
        },
    )
    return system


def test_markdown_ssp_covers_every_requirement(client, populated):
    response = client.get(f"/api/systems/{populated['id']}/ssp", params={"format": "md"})
    assert response.status_code == 200
    text = response.text
    for control in catalog.controls():
        assert control["cmmc_id"] in text, control["id"]
    assert "Precision Machining LLC" in text
    assert "Unique accounts in Entra ID." in text
    assert "No VoIP is used." in text
    assert "attachment;" in response.headers["content-disposition"]


def test_ssp_documents_scope_and_exclusions(client, populated):
    text = client.get(f"/api/systems/{populated['id']}/ssp", params={"format": "md"}).text
    assert "Specialized Assets" in text
    assert "Out-of-scope Assets" in text
    assert "Guest VLAN with no route to the enclave." in text
    assert "VLAN 20 behind the FortiGate." in text


def test_ssp_marks_undocumented_requirements_and_links_the_poam(client, populated):
    text = client.get(f"/api/systems/{populated['id']}/ssp", params={"format": "md"}).text
    assert "NOT YET DOCUMENTED" in text
    assert "POA&M-1" in text


def test_ssp_includes_assessment_objectives_and_evidence(client, populated):
    text = client.get(f"/api/systems/{populated['id']}/ssp", params={"format": "md"}).text
    assert "3.1.1[a]" in text
    assert "authorized users are identified" in text
    assert "Access Control Policy" in text
    assert "os.disk.encryption" in text


def test_ssp_html_renders_tables(client, populated):
    response = client.get(f"/api/systems/{populated['id']}/ssp", params={"format": "html"})
    assert response.status_code == 200
    assert response.text.startswith("<!doctype html>")
    assert "<table>" in response.text and "<h1>" in response.text
    assert "@media print" in response.text


def test_ssp_docx_opens_with_headings_and_tables(client, populated):
    from docx import Document

    response = client.get(f"/api/systems/{populated['id']}/ssp", params={"format": "docx"})
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX
    document = Document(io.BytesIO(response.content))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert any("System Security Plan" in h for h in headings)
    assert any("Access Control" in h for h in headings)
    assert len(document.tables) > 50


def test_ssp_json_carries_the_whole_context(client, populated):
    response = client.get(f"/api/systems/{populated['id']}/ssp", params={"format": "json"})
    context = json.loads(response.text)
    assert len(context["families"]) == 14
    assert sum(len(f["controls"]) for f in context["families"]) == 110
    assert context["score"]["total"] == 110


def test_poam_csv_has_the_assessor_columns(client, populated):
    import csv

    response = client.get(f"/api/systems/{populated['id']}/poam/export", params={"format": "csv"})
    assert response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert rows, "expected at least one POA&M row"
    row = next(r for r in rows if r["Security requirement"] == "3.5.3")
    assert row["CMMC practice"] == "IA.L2-3.5.3"
    assert row["Point value"] == "5"
    assert row["Weakness"] == "Roll out MFA"
    assert "Buy keys" in row["Milestones"]


def test_poam_markdown_and_docx(client, populated):
    markdown = client.get(f"/api/systems/{populated['id']}/poam/export", params={"format": "md"})
    assert "Plan of Action and Milestones" in markdown.text
    assert "POA&M-1" in markdown.text
    docx = client.get(f"/api/systems/{populated['id']}/poam/export", params={"format": "docx"})
    assert docx.headers["content-type"] == DOCX
    from docx import Document

    assert Document(io.BytesIO(docx.content)).tables


def test_readiness_report_states_the_score_and_blockers(client, populated):
    response = client.get(f"/api/systems/{populated['id']}/readiness-report", params={"format": "md"})
    text = response.text
    assert "CMMC Level 2 assessment readiness" in text
    assert "SPRS score" in text
    assert "Next actions" in text
    assert "Progress by family" in text


def test_readiness_report_flags_contradicted_requirements(client, populated):
    """The CNC controller reports unencrypted disk while 3.13.16 is claimed implemented."""
    sid = populated["id"]
    client.put(f"/api/systems/{sid}/controls/3.13.16", json={"status": "implemented",
                                                             "implementation_narrative": "BitLocker."})
    readiness = client.get(f"/api/systems/{sid}/readiness").json()
    assert readiness["contradiction_count"] >= 1
    assert any(c["control_id"] == "3.13.16" for c in readiness["contradictions"])
    text = client.get(f"/api/systems/{sid}/readiness-report", params={"format": "md"}).text
    assert "contradicted by live endpoint data" in text
    assert "os.disk.encryption" in text


def test_readiness_report_html_and_docx(client, populated):
    html = client.get(f"/api/systems/{populated['id']}/readiness-report", params={"format": "html"})
    assert html.text.startswith("<!doctype html>")
    docx = client.get(f"/api/systems/{populated['id']}/readiness-report", params={"format": "docx"})
    assert docx.headers["content-type"] == DOCX


def test_documents_require_authentication(anon_client, system):
    assert anon_client.get(f"/api/systems/{system['id']}/ssp").status_code == 401
