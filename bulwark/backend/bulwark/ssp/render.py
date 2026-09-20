"""Render the SSP context as Markdown, printable HTML or a Word document."""

from __future__ import annotations

import html
import io
import re
from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader, StrictUndefined

TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"

_PRINT_CSS = """
:root { color-scheme: light; }
body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
       max-width: 52rem; margin: 0 auto; padding: 2.5rem 1.5rem 6rem; line-height: 1.55;
       color: #16191d; background: #fff; }
h1 { font-size: 1.9rem; border-bottom: 2px solid #16191d; padding-bottom: .4rem; }
h2 { font-size: 1.4rem; margin-top: 2.4rem; border-bottom: 1px solid #d8dde3; padding-bottom: .3rem; }
h3 { font-size: 1.15rem; margin-top: 1.8rem; }
h4 { font-size: 1rem; margin-top: 1.5rem; color: #23408e; }
table { border-collapse: collapse; width: 100%; margin: .9rem 0; font-size: .88rem; }
th, td { border: 1px solid #d8dde3; padding: .38rem .55rem; text-align: left; vertical-align: top; }
th { background: #f4f6f9; font-weight: 600; }
code { background: #f4f6f9; padding: .1rem .3rem; border-radius: 3px; font-size: .88em; }
blockquote { border-left: 3px solid #23408e; margin: 1rem 0; padding: .3rem 0 .3rem 1rem;
             color: #46505c; background: #f8fafc; }
hr { border: 0; border-top: 1px solid #d8dde3; margin: 2rem 0; }
@media print { body { max-width: none; padding: 0; font-size: 10.5pt; }
                h2 { page-break-before: auto; } table { page-break-inside: avoid; } }
"""


def _environment() -> Environment:
    env = Environment(
        loader=FileSystemLoader(str(TEMPLATE_DIR)),
        undefined=StrictUndefined,
        trim_blocks=True,
        lstrip_blocks=False,
        autoescape=False,
    )
    return env


#: Sentinel template name: the context already carries rendered Markdown under ``_markdown``.
INLINE_TEMPLATE = "__inline__"


def render_markdown(context: dict[str, Any], template: str = "ssp.md.j2") -> str:
    """The canonical rendering; every other format derives from this context.

    Passing ``template=INLINE_TEMPLATE`` renders ``context["_markdown"]`` as-is, so the report
    exports can reuse the HTML and DOCX writers without a Jinja template of their own.
    """
    if template == INLINE_TEMPLATE:
        return str(context.get("_markdown", ""))
    return _environment().get_template(template).render(**context)


# --------------------------------------------------------------------------------------------
# Markdown -> HTML (small converter: headings, tables, lists, emphasis, blockquotes)
# --------------------------------------------------------------------------------------------
_INLINE = (
    (re.compile(r"\*\*(.+?)\*\*", re.S), r"<strong>\1</strong>"),
    (re.compile(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", re.S), r"<em>\1</em>"),
    (re.compile(r"(?<!_)_(?!\s)(.+?)(?<!\s)_(?!_)", re.S), r"<em>\1</em>"),
    (re.compile(r"`(.+?)`"), r"<code>\1</code>"),
    (re.compile(r"\[(.+?)\]\((.+?)\)"), r'<a href="\2">\1</a>'),
)


def _inline(text: str) -> str:
    out = html.escape(text, quote=False)
    for pattern, replacement in _INLINE:
        out = pattern.sub(replacement, out)
    return out


def _table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def markdown_to_html(text: str) -> str:
    """Convert the subset of Markdown the SSP template emits."""
    lines = text.splitlines()
    out: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            out.append(f"<h{level}>{_inline(stripped[level:].strip())}</h{level}>")
            index += 1
            continue

        if stripped in ("---", "***", "___"):
            out.append("<hr>")
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\s*\|[\s:|-]+\|\s*$", lines[index + 1]
        ):
            header = _table_row(stripped)
            index += 2
            body: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                body.append(_table_row(lines[index]))
                index += 1
            head = "".join(f"<th>{_inline(cell)}</th>" for cell in header)
            rows = "".join(
                "<tr>" + "".join(f"<td>{_inline(cell)}</td>" for cell in row) + "</tr>" for row in body
            )
            out.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
            continue

        if stripped.startswith("> "):
            block: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            out.append(f"<blockquote><p>{_inline(' '.join(block))}</p></blockquote>")
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        number = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet or number:
            tag = "ul" if bullet else "ol"
            pattern = r"^\s*[-*]\s+(.*)$" if bullet else r"^\s*\d+\.\s+(.*)$"
            items: list[str] = []
            while index < len(lines):
                match = re.match(pattern, lines[index])
                if not match:
                    break
                items.append(f"<li>{_inline(match.group(1))}</li>")
                index += 1
            out.append(f"<{tag}>{''.join(items)}</{tag}>")
            continue

        paragraph: list[str] = []
        while index < len(lines) and lines[index].strip() and not re.match(
            r"^\s*(#|\||>|[-*]\s|\d+\.\s|---\s*$)", lines[index]
        ):
            paragraph.append(lines[index].strip())
            index += 1
        if paragraph:
            out.append(f"<p>{_inline(' '.join(paragraph))}</p>")
        else:
            index += 1
    return "\n".join(out)


def render_html(context: dict[str, Any], title: str | None = None, template: str = "ssp.md.j2") -> str:
    """Printable, self-contained HTML."""
    body = markdown_to_html(render_markdown(context, template))
    heading = title or (
        f"System Security Plan - {context['system']['name']}"
        if "system" in context
        else "Bulwark report"
    )
    return (
        "<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        f"<title>{html.escape(heading)}</title>\n<style>{_PRINT_CSS}</style>\n</head>\n<body>\n"
        f"{body}\n</body>\n</html>\n"
    )


# --------------------------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------------------------
_BOLD = re.compile(r"\*\*(.+?)\*\*", re.S)


def _add_rich_text(paragraph, text: str) -> None:
    """Write text into a python-docx paragraph, honouring **bold** and stripping _emphasis_."""
    position = 0
    for match in _BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(_plain(text[position : match.start()]))
        paragraph.add_run(_plain(match.group(1))).bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(_plain(text[position:]))


def _plain(text: str) -> str:
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", text)
    text = re.sub(r"(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)", r"\1", text)
    return text


def render_docx(context: dict[str, Any], template: str = "ssp.md.j2") -> bytes:
    """A Word document with real heading styles and tables, built from the same Markdown."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    text = render_markdown(context, template)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading_text = _plain(stripped[level:].strip())
            if level == 2:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            document.add_heading(heading_text, level=min(level, 4))
            index += 1
            continue

        if stripped in ("---", "***", "___"):
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\s*\|[\s:|-]+\|\s*$", lines[index + 1]
        ):
            header = _table_row(stripped)
            index += 2
            body: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                body.append(_table_row(lines[index]))
                index += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, value in zip(table.rows[0].cells, header, strict=False):
                cell.text = ""
                run = cell.paragraphs[0].add_run(_plain(value))
                run.bold = True
            for row in body:
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=False):
                    cell.text = ""
                    _add_rich_text(cell.paragraphs[0], value)
            document.add_paragraph()
            continue

        if stripped.startswith(">"):
            block: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = document.add_paragraph(style="Intense Quote")
            _add_rich_text(paragraph, " ".join(block))
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        number = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet or number:
            pattern = r"^\s*[-*]\s+(.*)$" if bullet else r"^\s*\d+\.\s+(.*)$"
            list_style = "List Bullet" if bullet else "List Number"
            while index < len(lines):
                match = re.match(pattern, lines[index])
                if not match:
                    break
                paragraph = document.add_paragraph(style=list_style)
                _add_rich_text(paragraph, match.group(1))
                index += 1
            continue

        paragraph_lines: list[str] = []
        while index < len(lines) and lines[index].strip() and not re.match(
            r"^\s*(#|\||>|[-*]\s|\d+\.\s|---\s*$)", lines[index]
        ):
            paragraph_lines.append(lines[index].strip())
            index += 1
        if paragraph_lines:
            paragraph = document.add_paragraph()
            _add_rich_text(paragraph, " ".join(paragraph_lines))
        else:
            index += 1

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
